# GuiMenu.py
import os
import re
import shutil
import sys
import threading
import unicodedata
import tkinter as tk
from tkinter import filedialog, ttk, messagebox

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(THIS_DIR)
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)
IMPLEMENTACIONES_ROOT = r"Z:\IMPLEMENTACIONES"
AUDIT_EXCEL_PATH = os.path.join(BASE_DIR, "Plantillas", "PMO-ER-005.xlsx")
REPORT_TEMPLATE_PATH = os.path.join(BASE_DIR, "Plantillas", "PMO-ER-005.xlsx")

ALLOWED_SERVICES = (
    "FACTURACIÓN ELECTRONICA",
    "DOCUMENTO SOPORTE CONTABLE",
    "NOMINA ELECTRONICA",
    "RECEPCION",
    "DOCUMENTO EQUIVALENTE ELECTRÓNICO",
    "RADIAN",
    "RIPS",
)

SERVICE_FLOW_MODULES = {
    "facturacion electronica": "Flujos.Flujo_FE.FlujoAuditoriaFE",
    "documento soporte contable": "Flujos.Flujo_DS.FlujoAuditoriaDS",
    "nomina electronica": "Flujos.Flujo_NE.FlujoAuditoriaNE",
    "recepcion": "Flujos.Flujo_RE.FlujoAuditoriaRE",
    "documento equivalente electronico": "Flujos.Flujo_DE.FlujoAuditoriaDE",
    "radian": "Flujos.Flujo_RADIAN.FlujoAuditoriaRADIAN",
    "rips": "Flujos.Flujo_RIPS.FlujoAuditoriaRIPS",
}


def _get_service_flow(service_name):
    module_name = SERVICE_FLOW_MODULES.get(_normalize_text(service_name))
    if not module_name:
        raise ValueError(f"No existe un flujo configurado para el servicio '{service_name}'.")
    return __import__(module_name, fromlist=["*"])


def _normalize_text(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("&", "y")
    text = "".join(ch for ch in text if ch.isalnum() or ch.isspace())
    text = " ".join(text.split())
    return text

def _load_audit_tasks():
    if not os.path.isfile(AUDIT_EXCEL_PATH) or load_workbook is None:
        return []

    try:
        workbook = load_workbook(AUDIT_EXCEL_PATH, read_only=True, data_only=True)
        sheet_name = "PMO-ER-005" if "PMO-ER-005" in workbook.sheetnames else workbook.sheetnames[0]
        sheet = workbook[sheet_name]
        sections = []
        current_section = None
        for row in sheet.iter_rows(min_row=10, max_row=45, min_col=2, max_col=4, values_only=True):
            first_value, second_value, description = row
            first_text = str(first_value).strip() if first_value is not None else ""
            second_text = str(second_value).strip() if second_value is not None else ""
            if re.match(r"^\d+(\.\d+)?$", second_text):
                code_text, folder_text = second_text, first_text
            else:
                code_text, folder_text = first_text, second_text
            description_text = " ".join(str(description or "").split())
            if code_text.isdigit() and description_text:
                section_titles = {
                    "1": "COMERCIAL",
                    "2": "PRELIMINARES",
                    "3": "PRUEBAS",
                    "4": "CONFIGURACIÓN Y PARAMETRIZACIÓN",
                    "5": "SALIDA EN VIVO",
                    "6": "ACOMPAÑAMIENTO Y CIERRE",
                }
                current_section = {
                    "code": code_text,
                    "title": section_titles.get(code_text, description_text),
                    "folder": folder_text,
                    "items": [],
                }
                sections.append(current_section)
            elif re.match(r"^\d+\.\d+$", code_text) and description_text and current_section is not None:
                current_section["items"].append({
                    "code": code_text,
                    "title": description_text,
                    "folder": folder_text,
                })
        workbook.close()
        return sections
    except Exception as exc:
        return [{"code": "", "title": f"No se pudo leer la auditoría desde Excel: {exc}", "folder": "", "items": []}]


def _match_client_score(folder_name, client_input):
    client_input = _normalize_text(client_input)
    folder_name = _normalize_text(folder_name)

    if not client_input:
        return 0

    if folder_name == client_input:
        return 100
    if client_input in folder_name:
        return 75 + min(20, len(client_input))

    client_tokens = [token for token in client_input.split() if token]
    matching_tokens = sum(1 for token in client_tokens if token in folder_name)
    if matching_tokens:
        return 50 + min(25, matching_tokens * 10)

    return 0


def _list_implementation_folders():
    if not os.path.isdir(IMPLEMENTACIONES_ROOT):
        return []

    allowed_services = {_normalize_text(service) for service in ALLOWED_SERVICES}
    folders = []
    for name in sorted(os.listdir(IMPLEMENTACIONES_ROOT)):
        full_path = os.path.join(IMPLEMENTACIONES_ROOT, name)
        if os.path.isdir(full_path) and _normalize_text(name) in allowed_services:
            folders.append(name)

    return sorted(folders, key=_normalize_text)


def _list_service_clients(service_name):
    if not service_name:
        return []

    service_dir = os.path.join(IMPLEMENTACIONES_ROOT, service_name)
    if not os.path.isdir(service_dir):
        return []

    clients = []
    for name in sorted(os.listdir(service_dir)):
        full_path = os.path.join(service_dir, name)
        if os.path.isdir(full_path):
            clients.append(name)
    return clients


def _find_matching_client_in_service(service_name, client_name):
    service_dir = os.path.join(IMPLEMENTACIONES_ROOT, service_name)
    if not os.path.isdir(service_dir):
        if _match_client_score(service_name, client_name) >= 40:
            return service_dir
        return None

    candidates = []
    for name in sorted(os.listdir(service_dir)):
        full_path = os.path.join(service_dir, name)
        if os.path.isdir(full_path):
            candidates.append((name, full_path))

    if not candidates:
        if _match_client_score(service_name, client_name) >= 40:
            return service_dir
        return None

    best_match = None
    best_score = -1
    for folder_name, full_path in candidates:
        score = _match_client_score(folder_name, client_name)
        if score > best_score:
            best_score = score
            best_match = full_path

    if best_match is not None and best_score >= 40:
        return best_match

    if _match_client_score(service_name, client_name) >= 40:
        return service_dir

    return None


class AuditCancelled(Exception):
    pass


class MissingEvidenceDialog(tk.Toplevel):
    def __init__(self, parent, requirement, validate_requirement_file):
        super().__init__(parent)
        self.title("Evidencia faltante")
        self.requirement = requirement
        self.validate_requirement_file = validate_requirement_file
        self.result = None
        self.geometry("620x350")
        self.transient(parent)
        self.grab_set()

        tk.Label(self, text=f"Requisito: {requirement['id']} - {requirement.get('name', '')}", anchor="w", justify="left").pack(fill="x", padx=16, pady=(12, 6))
        tk.Label(self, text=f"Carpeta esperada: {requirement.get('folder', '')}", anchor="w").pack(fill="x", padx=16, pady=(0, 12))

        self.file_var = tk.StringVar()
        btn_row = tk.Frame(self)
        btn_row.pack(fill="x", padx=16, pady=(0, 8))
        tk.Button(btn_row, text="Cargar archivo/correo", command=self._browse_file).pack(side="left")
        tk.Entry(btn_row, textvariable=self.file_var, width=60).pack(side="left", padx=(8, 0), fill="x", expand=True)

        self.not_applicable_var = tk.BooleanVar()
        self.not_applicable_checkbox = tk.Checkbutton(self, text="¿No aplica para este cliente?", variable=self.not_applicable_var, command=self._toggle_na)
        self.not_applicable_checkbox.pack(anchor="w", padx=16, pady=(6, 8))

        self.na_frame = tk.LabelFrame(self, text="Justificación para N/A")
        self.na_frame.pack(fill="both", expand=True, padx=16, pady=(0, 8))
        self.na_text = tk.Text(self.na_frame, height=6)
        self.na_text.pack(fill="both", expand=True, padx=8, pady=8)
        self.na_frame.pack_forget()

        button_row = tk.Frame(self)
        button_row.pack(fill="x", padx=16, pady=(0, 12))
        tk.Button(button_row, text="Cancelar auditoría", command=self._cancel_audit).pack(side="left")
        tk.Button(button_row, text="Guardar", command=self._save).pack(side="right")
        tk.Button(button_row, text="Cancelar", command=self.destroy).pack(side="right", padx=(0, 8))

        self.update_idletasks()
        self._center_on_parent(parent)
        self.lift()
        self.focus_force()

    def _center_on_parent(self, parent):
        width = self.winfo_width()
        height = self.winfo_height()
        parent_width = parent.winfo_width()
        parent_height = parent.winfo_height()
        parent_x = parent.winfo_rootx()
        parent_y = parent.winfo_rooty()
        x = parent_x + max(0, (parent_width - width) // 2)
        y = parent_y + max(0, (parent_height - height) // 2)
        self.geometry(f"{width}x{height}+{x}+{y}")

    def _toggle_na(self):
        if self.not_applicable_var.get():
            self.na_frame.pack(fill="both", expand=True, padx=16, pady=(0, 8))
        else:
            self.na_frame.pack_forget()

    def _browse_file(self):
        path = filedialog.askopenfilename(title="Seleccionar evidencia", filetypes=[("Todos los archivos", "*.*")])
        if path:
            self.file_var.set(path)

    def _cancel_audit(self):
        self.result = {"cancel_audit": True}
        self.destroy()

    def _save(self):
        file_path = self.file_var.get().strip()
        not_applicable = bool(self.not_applicable_var.get())
        observation = self.na_text.get("1.0", "end").strip() if not_applicable else ""

        if not_applicable:
            if not observation:
                messagebox.showwarning("Falta justificación", "Debe escribir el motivo por el cual no aplica para este cliente.")
                return
            self.result = {"not_applicable": True, "observation": observation}
        elif file_path:
            ok, error = self.validate_requirement_file(file_path, self.requirement)
            if not ok:
                messagebox.showwarning("Código o nombre inválido", error)
                return
            self.result = {"path": file_path, "observation": f"Archivo cargado manualmente: {file_path}"}
        else:
            messagebox.showwarning("Falta evidencia", "Debe cargar un archivo o marcar la opción 'No aplica'.")
            return

        self.destroy()


def _store_manual_evidence(client_path, requirement, source_path):
    folder = requirement.get("folder", "").replace("/", os.sep).replace("\\", os.sep)
    destination_dir = os.path.join(client_path, folder)
    os.makedirs(destination_dir, exist_ok=True)
    destination_path = os.path.join(destination_dir, os.path.basename(source_path))
    if os.path.abspath(source_path) != os.path.abspath(destination_path):
        shutil.copy2(source_path, destination_path)
    return destination_path


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("A.I.D")
        self.geometry("1200x720")
        self.configure(bg="#FFFFFF")
        self.client_path = None
        self.available_client_folders = _list_implementation_folders()
        self.audit_tasks = _load_audit_tasks()
        self._build_style()
        self._build_ui()

    def _build_style(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure(".", background="#FFFFFF", foreground="#111111", fieldbackground="#FFFFFF")
        style.configure("TLabel", font=("Segoe UI", 10), foreground="#111111")
        style.configure("Header.TLabel", font=("Segoe UI", 16, "bold"), foreground="#111111")
        style.configure("Card.TLabelframe", background="#FFFFFF", foreground="#111111")
        style.configure("Card.TLabelframe.Label", font=("Segoe UI", 11, "bold"), foreground="#111111")
        style.configure("TButton", font=("Segoe UI", 10), padding=7, foreground="#111111", background="#FFFFFF")
        style.map("TButton", foreground=[("disabled", "#6B6B6B")], background=[("active", "#EAEAEA"), ("!disabled", "#FFFFFF")])
        style.configure("Audit.Treeview", background="#FFFFFF", fieldbackground="#FFFFFF", foreground="#111111", rowheight=28, font=("Segoe UI", 10))
        style.configure("Audit.Treeview.Heading", background="#F1F4F8", foreground="#111111", font=("Segoe UI", 10, "bold"))
        style.map("Audit.Treeview", background=[("selected", "#D9E8FF")], foreground=[("selected", "#111111")])

    def _build_ui(self):
        ttk.Label(self, text="A.I.D", style="Header.TLabel").pack(anchor="w", padx=16, pady=(16, 8))

        main = ttk.Frame(self)
        main.pack(fill="both", expand=True, padx=16, pady=(0, 8))

        left = ttk.Labelframe(main, text="Cliente", style="Card.TLabelframe")
        left.pack(side="left", fill="both", expand=True, padx=(0, 8))

        form = ttk.Frame(left)
        form.pack(fill="x", padx=12, pady=(12, 8))

        ttk.Label(form, text="Servicio:").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=(0, 6))
        self.folder_var = tk.StringVar()
        folder_values = self.available_client_folders if self.available_client_folders else ["Sin carpetas disponibles"]
        self.folder_combo = ttk.Combobox(form, textvariable=self.folder_var, values=folder_values, state="readonly" if self.available_client_folders else "disabled", width=32)
        self.folder_combo.grid(row=0, column=1, sticky="ew", pady=(0, 6))

        ttk.Label(form, text="Cliente:").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=(0, 6))
        self.client_var = tk.StringVar()
        self.client_options = []
        self._updating_client_options = False
        self.client_combo = ttk.Combobox(form, textvariable=self.client_var, values=[], width=32)
        self.client_combo.grid(row=1, column=1, sticky="ew", pady=(0, 6))
        self.client_combo.bind("<Return>", lambda event: self._validate_client())
        self.client_combo.bind("<<ComboboxSelected>>", self._on_client_selected)
        self.client_var.trace_add("write", self._on_client_changed)

        self.client_matches_list = tk.Listbox(form, height=5, exportselection=False, takefocus=0)
        self.client_matches_list.grid(row=2, column=1, sticky="ew", pady=(0, 6))
        self.client_matches_list.grid_remove()
        self.client_matches_list.bind("<Button-1>", self._select_client_match)

        form.columnconfigure(1, weight=1)

        buttons_row = ttk.Frame(left)
        buttons_row.pack(fill="x", padx=12, pady=(0, 12))
        self.start_button = ttk.Button(buttons_row, text="Iniciar auditoría", command=self._start_audit, state="disabled")
        self.start_button.pack(side="left", padx=(0, 8))
        ttk.Button(buttons_row, text="Cancelar auditoría", command=self._cancel_audit).pack(side="left")

        self.folder_var.trace_add("write", self._on_service_changed)
        self._refresh_service_clients()

        tasks_frame = ttk.Labelframe(left, text="Secciones de auditoría", style="Card.TLabelframe")
        tasks_frame.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self.tasks_toggle = ttk.Button(tasks_frame, text="Minimizar", command=self._toggle_audit_tree)
        self.tasks_toggle.pack(anchor="e", padx=8, pady=(8, 0))
        self.tree_container = ttk.Frame(tasks_frame)
        self.tree_container.pack(fill="both", expand=True, padx=8, pady=(4, 8))
        self.audit_tree = ttk.Treeview(
            self.tree_container,
            columns=("folder", "requirement"),
            show="tree headings",
            style="Audit.Treeview",
            selectmode="browse",
        )
        self.audit_tree.heading("#0", text="Sección / subsección", anchor="w")
        self.audit_tree.heading("folder", text="Carpeta", anchor="w")
        self.audit_tree.heading("requirement", text="Documento o evidencia requerida", anchor="w")
        self.audit_tree.column("#0", width=190, minwidth=150, stretch=False)
        self.audit_tree.column("folder", width=130, minwidth=100, stretch=False)
        self.audit_tree.column("requirement", width=470, minwidth=250, stretch=True)
        scrollbar = ttk.Scrollbar(self.tree_container, orient="vertical", command=self.audit_tree.yview)
        self.audit_tree.configure(yscrollcommand=scrollbar.set)
        self.audit_tree.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self._populate_audit_tree()
        self._audit_tree_expanded = True

        log_box = ttk.Labelframe(self, text="Salida", style="Card.TLabelframe")
        log_box.pack(fill="both", expand=False, padx=16, pady=(0, 16))
        self.txt = tk.Text(log_box, height=10, bg="#FFFFFF", fg="#111111", insertbackground="#111111", font=("Consolas", 10))
        self.txt.pack(fill="both", expand=True, padx=8, pady=8)

        self._log("Sistema listo.")
        self._log(f"Ruta compartida: {IMPLEMENTACIONES_ROOT}")
        self._log(f"Archivo de auditoría: {AUDIT_EXCEL_PATH}")

    def _populate_audit_tree(self):
        for section in self.audit_tasks:
            section_id = self.audit_tree.insert(
                "",
                "end",
                text=f"{section['code']}. {section['title']}",
                values=(section.get("folder", ""), ""),
                tags=("section",),
                open=True,
            )
            for item in section.get("items", []):
                self.audit_tree.insert(
                    section_id,
                    "end",
                    text=item["code"],
                    values=(item.get("folder", ""), item["title"]),
                    tags=("subsection",),
                )
        self.audit_tree.tag_configure("section", background="#EAF0F7", font=("Segoe UI", 10, "bold"))
        self.audit_tree.tag_configure("subsection", background="#FFFFFF")

    def _toggle_audit_tree(self):
        self._audit_tree_expanded = not getattr(self, "_audit_tree_expanded", True)
        if self._audit_tree_expanded:
            self.tree_container.pack(fill="both", expand=True, padx=8, pady=(4, 8))
            self.tasks_toggle.configure(text="Minimizar")
        else:
            self.tree_container.pack_forget()
            self.tasks_toggle.configure(text="Expandir")

    def _cancel_audit(self):
        self._set_start_enabled(False)
        self.client_path = None
        self.client_var.set("")
        self._log("Auditoría cancelada por el usuario.")
        messagebox.showinfo("Auditoría cancelada", "La auditoría fue cancelada y se limpió la selección del cliente.")

    def _log(self, message):
        self.txt.insert("end", str(message) + "\n")
        self.txt.see("end")

    def _set_start_enabled(self, enabled):
        self.start_button.configure(state="normal" if enabled else "disabled")

    def _on_service_changed(self, *args):
        self._refresh_service_clients()

    def _on_client_changed(self, *args):
        if self._updating_client_options:
            return

        self.client_path = None
        self._set_start_enabled(False)
        client_input = self.client_var.get().strip()
        if not client_input:
            self.client_combo.configure(values=self.client_options)
            self.client_matches_list.grid_remove()
            return

        matches = [
            (client_name, _match_client_score(client_name, client_input))
            for client_name in self.client_options
        ]
        matches = [match for match in matches if match[1] >= 40]
        matches.sort(key=lambda match: (-match[1], _normalize_text(match[0])))
        filtered_clients = [client_name for client_name, _ in matches]

        self._updating_client_options = True
        try:
            self.client_combo.configure(values=filtered_clients)
        finally:
            self._updating_client_options = False

        self.client_matches_list.delete(0, "end")
        for client_name in filtered_clients:
            self.client_matches_list.insert("end", client_name)
        if filtered_clients:
            self.client_matches_list.grid()
        else:
            self.client_matches_list.grid_remove()

    def _select_client_match(self, event):
        selected_index = self.client_matches_list.nearest(event.y)
        if selected_index < 0 or selected_index >= self.client_matches_list.size():
            return "break"

        selected_client = self.client_matches_list.get(selected_index)
        self._updating_client_options = True
        try:
            self.client_var.set(selected_client)
        finally:
            self._updating_client_options = False
        self.client_matches_list.grid_remove()
        self._validate_client()
        return "break"

    def _on_client_selected(self, event=None):
        self.after_idle(self._validate_client)

    def _refresh_service_clients(self):
        selected_service = self.folder_var.get().strip()
        clients = _list_service_clients(selected_service) if selected_service else []
        self.client_options = clients
        self._updating_client_options = True
        try:
            self.client_combo.configure(values=clients)
            self.client_var.set("")
        finally:
            self._updating_client_options = False
        self.client_matches_list.grid_remove()
        self.client_combo.configure(state="normal" if clients else "disabled")
        self.client_path = None
        self._set_start_enabled(False)

    def _validate_client(self):
        client_name = self.client_var.get().strip()
        selected_folder = self.folder_var.get().strip()

        if not os.path.isdir(IMPLEMENTACIONES_ROOT):
            messagebox.showerror("Ruta compartida no disponible", f"No se encontró la ruta compartida: {IMPLEMENTACIONES_ROOT}")
            self.client_path = None
            self._set_start_enabled(False)
            self._log(f"No se pudo validar cliente: {IMPLEMENTACIONES_ROOT} no está disponible.")
            return

        if not selected_folder:
            messagebox.showwarning("Carpeta requerida", "Seleccione una carpeta de la lista antes de continuar.")
            self.client_path = None
            self._set_start_enabled(False)
            return

        if not client_name:
            messagebox.showwarning("Cliente requerido", "Debe escribir o seleccionar el nombre del cliente para validarlo contra la carpeta seleccionada.")
            self.client_path = None
            self._set_start_enabled(False)
            return

        service_dir = os.path.join(IMPLEMENTACIONES_ROOT, selected_folder)
        if not os.path.isdir(service_dir):
            messagebox.showerror("Carpeta inexistente", f"La carpeta seleccionada no existe en {IMPLEMENTACIONES_ROOT}.")
            self.client_path = None
            self._set_start_enabled(False)
            self._log(f"Carpeta no encontrada: {service_dir}")
            return

        folder_path = _find_matching_client_in_service(selected_folder, client_name)
        if folder_path is None:
            self.client_path = None
            self._set_start_enabled(False)
            messagebox.showerror("Validación fallida", f"El cliente '{client_name}' no coincide con ninguna carpeta del servicio '{selected_folder}'.")
            self._log(f"Validación fallida: cliente '{client_name}' vs servicio '{selected_folder}'")
            return

        self.client_path = folder_path
        self._set_start_enabled(True)
        self._log(f"Cliente validado: {os.path.basename(folder_path)}")
        self._log(f"Ruta válida: {folder_path}")

    def _start_audit(self):
        if not self.client_path:
            messagebox.showwarning("Cliente no validado", "Primero debe validar el cliente para iniciar la auditoría.")
            return

        selected_service = self.folder_var.get().strip()
        try:
            service_flow = _get_service_flow(selected_service)
        except ValueError as exc:
            messagebox.showerror("Flujo no disponible", str(exc))
            return

        client_path = self.client_path
        self._set_start_enabled(False)
        self._log(f"Iniciando auditoría de {service_flow.SERVICE_NAME} en segundo plano para: {client_path}")
        threading.Thread(
            target=self._run_audit_in_background,
            args=(client_path, service_flow),
            daemon=True,
        ).start()

    def _prompt_missing_requirement(self, requirement, service_flow):
        dialog = MissingEvidenceDialog(self, requirement, service_flow.validate_uploaded_requirement_file)
        self.wait_window(dialog)
        return dialog.result

    def _collect_manual_overrides(self, client_path, service_flow):
        manual_overrides = {}
        for requirement in service_flow.REQUIREMENT_CATALOG:
            requirement_id = requirement.get("id")
            propagated = service_flow._propagate_not_applicable_override(manual_overrides, requirement_id)
            if propagated:
                manual_overrides[requirement_id] = propagated
                continue

            result = service_flow.evaluate_requirement(client_path, requirement)
            if result.get("status") != "NO":
                continue
            if result.get("path") and requirement_id != "4.1":
                continue
            if requirement_id == "4.1":
                loaded_paths = []
                for evidence_type in ("archivo", "correo"):
                    prompt_requirement = dict(requirement)
                    prompt_requirement["name"] = f"{requirement.get('name', '')} ({evidence_type})"
                    response = self._prompt_missing_requirement(prompt_requirement, service_flow)
                    if response is None:
                        break
                    if response.get("cancel_audit"):
                        raise AuditCancelled()
                    if response.get("not_applicable"):
                        manual_overrides[requirement_id] = {
                            "not_applicable": True,
                            "observation": response.get("observation") or "No aplica para este cliente.",
                        }
                        break
                    if response.get("path"):
                        loaded_paths.append(_store_manual_evidence(client_path, requirement, response["path"]))
                if len(loaded_paths) == 2 and requirement_id not in manual_overrides:
                    manual_overrides[requirement_id] = {"paths": loaded_paths}
                continue
            response = self._prompt_missing_requirement(requirement, service_flow)
            if response is None:
                continue
            if response.get("cancel_audit"):
                raise AuditCancelled()
            if response.get("not_applicable"):
                manual_overrides[requirement_id] = {
                    "not_applicable": True,
                    "observation": response.get("observation") or "No aplica para este cliente.",
                }
            elif response.get("path"):
                stored_path = _store_manual_evidence(client_path, requirement, response["path"])
                manual_overrides[requirement_id] = {
                    "not_applicable": False,
                    "path": stored_path,
                    "observation": f"Documento encontrado en: {stored_path}",
                }
        return manual_overrides

    def _audit_finished(self, report_path, error):
        self._set_start_enabled(bool(self.client_path))
        if error:
            self._log(f"La auditoría terminó con error: {error}")
            messagebox.showerror("Auditoría no completada", error)
            return
        self._log(f"Auditoría completada. Reporte Excel: {report_path}")
        messagebox.showinfo("Auditoría completada", "La auditoría terminó correctamente y el reporte Excel fue guardado en la carpeta SALIDA EN VIVO del cliente.")

    def _audit_cancelled(self):
        self.client_path = None
        self.client_var.set("")
        self._set_start_enabled(False)
        self._log("Auditoría cancelada por el usuario.")
        messagebox.showinfo("Auditoría cancelada", "La auditoría fue cancelada y no se generó ningún reporte.")

    def _run_audit_in_background(self, client_path, service_flow):
        try:
            manual_overrides = self._collect_manual_overrides(client_path, service_flow)
            report_path = service_flow.run_audit(
                client_path,
                service_name=service_flow.SERVICE_NAME,
                template_path=REPORT_TEMPLATE_PATH,
                manual_overrides=manual_overrides,
            )
        except AuditCancelled:
            self.after(0, self._audit_cancelled)
            return
        except Exception as exc:
            self.after(0, lambda exc=exc: self._audit_finished(None, str(exc)))
            return
        self.after(0, lambda: self._audit_finished(report_path, None))


def run_gui():
    App().mainloop()


if __name__ == "__main__":
    run_gui()